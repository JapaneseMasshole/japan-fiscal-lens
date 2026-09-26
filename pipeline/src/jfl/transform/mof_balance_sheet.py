"""Parse the national balance sheet (貸借対照表) from MOF 国の財務書類 workbooks.

Sheet layout (stable FY2020–FY2023):
    rows 1-5   title, unit (単位：百万円), column headers 前会計年度 / 本会計年度
    left block   A=label  B=previous year  C=current year   (assets, ＜資産の部＞)
    right block  D=label  E=previous year  F=current year   (liabilities, then net assets)
    footnotes    rows starting with "(注"

Nesting is encoded with the Excel cell indent (0, 1, 2). 貸倒引当金 (allowance
for bad debts) is indented but is a top-level deduction from assets, not a
child of the line above it.

Each workbook holds two fiscal-year-ends: the file's own year ("current") and
the year before ("previous"). We keep both so they can be cross-checked.
"""

from __future__ import annotations

import warnings
from dataclasses import asdict, dataclass
from pathlib import Path

import openpyxl
import pandas as pd

from jfl.transform.excel import UnreadableWorkbook, check_readable_xlsx, norm

SHEET = "貸借対照表"
CONTRA_ITEMS = {"貸倒引当金"}
TOTALS = {
    "資産合計": "total_assets",
    "負債合計": "total_liabilities",
    "負債及び資産・負債差額合計": "total_liabilities_and_net_assets",
}
SECTION_MARKERS = {
    "＜資産の部＞": "assets",
    "＜負債の部＞": "liabilities",
    "＜資産・負債差額の部＞": "net_assets",
}
UNIT_DIVISORS_TO_OKU = {"百万円": 100, "千円": 100_000, "億円": 1}

ITEM_EN = {
    "現金・預金": "Cash and deposits",
    "有価証券": "Securities",
    "たな卸資産": "Inventories",
    "未収金": "Accounts receivable",
    "未収収益": "Accrued revenue",
    "未収（再）保険料": "(Re)insurance premiums receivable",
    "前払費用": "Prepaid expenses",
    "貸付金": "Loans",
    "運用寄託金": "Pension reserve deposits (GPIF)",
    "その他の債権等": "Other receivables",
    "貸倒引当金": "Allowance for bad debts",
    "有形固定資産": "Tangible fixed assets",
    "国有財産（公共用財産を除く）": "State property (excl. public-use)",
    "土地": "Land",
    "立木竹": "Standing timber",
    "建物": "Buildings",
    "工作物": "Structures",
    "機械器具": "Machinery",
    "船舶": "Vessels",
    "航空機": "Aircraft",
    "建設仮勘定": "Construction in progress",
    "公共用財産": "Public-use property (roads, rivers, ports…)",
    "公共用財産用地": "Public-use land",
    "公共用財産施設": "Public-use facilities",
    "物品": "Equipment",
    "その他の固定資産": "Other fixed assets",
    "無形固定資産": "Intangible fixed assets",
    "出資金": "Capital investments",
    "資産合計": "Total assets",
    "未払金": "Accounts payable",
    "支払備金": "Claims reserve",
    "未払費用": "Accrued expenses",
    "保管金等": "Deposits held",
    "前受金": "Advances received",
    "前受収益": "Unearned revenue",
    "未経過（再）保険料": "Unearned (re)insurance premiums",
    "賞与引当金": "Bonus reserve",
    "政府短期証券": "Financing bills (short-term)",
    "公債": "Government bonds",
    "借入金": "Borrowings",
    "預託金": "Deposits (FILP)",
    "責任準備金": "Policy reserves",
    "公的年金預り金": "Public pension reserves",
    "退職給付引当金": "Retirement benefit reserve",
    "その他の債務等": "Other liabilities",
    "負債合計": "Total liabilities",
    "資産・負債差額": "Net assets (assets minus liabilities)",
    "負債及び資産・負債差額合計": "Total liabilities and net assets",
}


@dataclass
class Row:
    fiscal_year: int  # fiscal year whose end the figure describes (FY2023 = 2024-03-31)
    column: str  # "current" or "previous" within the source workbook
    source_file_year: int
    section: str
    order: int
    item_ja: str
    item_en: str
    level: int
    parent_ja: str
    is_total: bool
    value_oku_yen: float


ERA_BASE = {"令和": 2018, "平成": 1988}
KANJI_DIGITS = str.maketrans("０１２３４５６７８９", "0123456789")


def _era_year(text: str) -> int | None:
    """Western year of a Japanese-era date.

    '(令和５年' -> 2023, '(平成31年' -> 2019, '(令和元年' -> 2019, '(至 令和6年3月31日)' -> 2024.
    """
    t = norm(text).translate(KANJI_DIGITS).replace("元", "1")
    for era, base in ERA_BASE.items():
        if era in t:
            year_part = t.split(era, 1)[1].split("年", 1)[0]
            digits = "".join(ch for ch in year_part if ch.isdigit())
            return base + int(digits) if digits else None
    return None


def _fiscal_years(ws, label_col: int, first_row: int) -> tuple[int, int]:
    """Read (previous FY, current FY) from the sheet's own 3月31日 date headers.

    A balance sheet dated 令和6年3月31日 closes fiscal year 2023. We trust these
    dates over file or folder names, because published links are sometimes wrong.
    """
    for r in range(1, first_row):
        prev_y = _era_year(ws.cell(r, label_col + 1).value)
        cur_y = _era_year(ws.cell(r, label_col + 2).value)
        if prev_y and cur_y:
            if cur_y != prev_y + 1:
                raise ValueError(f"{ws.title}: header dates {prev_y}/{cur_y} are not consecutive")
            return prev_y - 1, cur_y - 1
    raise ValueError(f"{ws.title}: could not read the 3月31日 date headers")


def _find_layout(ws) -> tuple[int, list[int], int]:
    """Return (unit divisor, label columns, first data row)."""
    divisor = None
    label_cols: list[int] = []
    first_row = None
    for row in ws.iter_rows(min_row=1, max_row=15):
        for cell in row:
            text = norm(cell.value)
            if "単位" in text:
                for unit, d in UNIT_DIVISORS_TO_OKU.items():
                    if unit in text:
                        divisor = d
            if text in ("＜資産の部＞", "＜負債の部＞"):
                label_cols.append(cell.column)
                first_row = cell.row
    if divisor is None:
        raise ValueError(f"{ws.title}: unit row (単位) not found")
    if sorted(label_cols) != label_cols or len(label_cols) != 2 or first_row is None:
        raise ValueError(f"{ws.title}: could not find ＜資産の部＞ / ＜負債の部＞ headers")
    return divisor, label_cols, first_row


def parse_workbook(path: Path, expected_year: int | None = None) -> list[Row]:
    check_readable_xlsx(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        raise ValueError(f"{path}: sheet {SHEET!r} not found (have {wb.sheetnames})")
    ws = wb[SHEET]
    divisor, label_cols, first_row = _find_layout(ws)
    prev_fy, file_year = _fiscal_years(ws, label_cols[0], first_row)
    if expected_year is not None and file_year != expected_year:
        warnings.warn(
            f"{path}: stored as FY{expected_year} but the sheet is dated FY{file_year}; "
            f"using FY{file_year}.",
            stacklevel=2,
        )

    rows: list[Row] = []
    for label_col in label_cols:
        section = None
        stack: list[tuple[int, str]] = []  # (indent, item) ancestors
        order = 0
        for r in range(first_row, ws.max_row + 1):
            cell = ws.cell(r, label_col)
            label = norm(cell.value)
            if not label:
                continue
            if label.startswith("(注") or label.startswith("（注"):
                break
            if label in SECTION_MARKERS:
                section = SECTION_MARKERS[label]
                stack = []
                continue
            prev_v = ws.cell(r, label_col + 1).value
            cur_v = ws.cell(r, label_col + 2).value
            if not isinstance(cur_v, int | float) or not isinstance(prev_v, int | float):
                raise ValueError(f"{path} {cell.coordinate} {label!r}: non-numeric values")

            indent = int(cell.alignment.indent or 0)
            if label in CONTRA_ITEMS or label in TOTALS:
                indent = 0
            while stack and stack[-1][0] >= indent:
                stack.pop()
            parent = stack[-1][1] if stack else ""
            stack.append((indent, label))
            order += 1

            for column, fy, value in (
                ("current", file_year, cur_v),
                ("previous", prev_fy, prev_v),
            ):
                rows.append(
                    Row(
                        fiscal_year=fy,
                        column=column,
                        source_file_year=file_year,
                        section="total" if label in TOTALS else section,
                        order=order,
                        item_ja=label,
                        item_en=ITEM_EN.get(label, ""),
                        level=indent,
                        parent_ja=parent,
                        is_total=label in TOTALS,
                        value_oku_yen=value / divisor,
                    )
                )
    missing_en = sorted({r.item_ja for r in rows if not r.item_en})
    if missing_en:
        warnings.warn(f"{path}: no English label for {missing_en}", stacklevel=2)
    return rows


def parse_all(raw_dir: Path, scope: str = "gassan") -> tuple[pd.DataFrame, list[str]]:
    """Parse every readable workbook. Returns (all rows, list of skipped-file messages)."""
    rows: list[dict] = []
    skipped: list[str] = []
    for year_dir in sorted(raw_dir.glob("fy*")):
        path = year_dir / f"{scope}.xlsx"
        if not path.exists():
            continue
        try:
            parsed = [asdict(r) for r in parse_workbook(path, int(year_dir.name[2:]))]
        except UnreadableWorkbook as e:
            skipped.append(str(e))
            continue
        for r in parsed:
            r["source_folder"] = year_dir.name
        rows += parsed
    df = pd.DataFrame(rows)
    df.insert(0, "scope", scope)
    return df, skipped


def canonical(all_rows: pd.DataFrame) -> pd.DataFrame:
    """One value per (fiscal_year, item): prefer the year's own workbook ("current").

    A "previous" column is used only for years whose own workbook is missing or
    unreadable (e.g. the year before the earliest file we have).
    """
    df = all_rows.copy()
    # Prefer "current"; among several, prefer the newest workbook (it may carry restatements).
    df["_rank"] = (df["column"] != "current").astype(int)
    df = df.sort_values(
        ["fiscal_year", "section", "order", "_rank", "source_file_year"],
        ascending=[True, True, True, True, False],
    )
    key = ["fiscal_year", "section", "order", "item_ja"]
    out = df.drop_duplicates(key, keep="first").drop(columns="_rank")
    return out.reset_index(drop=True)
